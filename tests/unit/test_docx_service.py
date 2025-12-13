"""
Testes Unitários - DOCX Service

Testa o serviço de manipulação de documentos DOCX em app/services/docx_service.py

Autor: Maxwell da Silva Oliveira - M&S do Brasil LTDA
Versão: 1.5.2
"""
import pytest
import io
from docx import Document
from docx.document import Document as DocumentType
from app.services.docx_service import DocxService


class TestReplaceTagsInDoc:
    """Testes para DocxService.replace_tags_in_doc()"""

    def test_replace_tags_lista_vendas(self, lista_vendas_bytes, replacements_lista_vendas):
        """Deve substituir tags em ListaVendas.docx"""
        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements_lista_vendas
        )

        assert isinstance(result_doc, DocumentType)

        # Extrai todo o texto do documento
        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Verifica se as substituições foram feitas
        assert 'João Silva' in text
        assert '06/12/2025' in text
        assert 'R$ 1.500,00' in text

        # Verifica se as tags foram removidas
        assert '{NOME_CLIENTE}' not in text
        assert '{DATA_VENDA}' not in text
        assert '{VALOR}' not in text

    def test_replace_tags_with_tags_doc(self, with_tags_bytes, replacements_with_tags):
        """Deve substituir tags em with_tags.docx"""
        result_doc = DocxService.replace_tags_in_doc(
            with_tags_bytes,
            replacements_with_tags
        )

        # Extrai texto
        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Verifica substituições
        assert 'Dezembro' in text
        assert '2025' in text
        assert 'M&S do Brasil LTDA' in text
        assert 'Maxwell da Silva Oliveira' in text

        # Verifica remoção de tags
        assert '{MES}' not in text
        assert '{ANO}' not in text
        assert '{EMPRESA}' not in text

    def test_replace_tags_complex_doc(self, complex_doc_bytes, replacements_complex):
        """Deve substituir tags em documento complexo"""
        result_doc = DocxService.replace_tags_in_doc(
            complex_doc_bytes,
            replacements_complex
        )

        # Extrai texto dos parágrafos
        paragraphs_text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Verifica substituições em parágrafos
        assert 'Empresa XYZ' in paragraphs_text
        assert 'Implementação API' in paragraphs_text

        # Extrai texto das tabelas
        tables_text = ''
        for table in result_doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    tables_text += cell.text + ' '

        # Verifica substituições em tabelas
        assert 'Desenvolvimento' in tables_text
        assert '160h' in tables_text
        assert 'R$ 16.000,00' in tables_text
        assert 'R$ 20.000,00' in tables_text

    def test_replace_tags_repetidas(self, tags_repetidas_bytes, replacements_tags_repetidas):
        """Deve substituir todas as ocorrências de tags repetidas"""
        result_doc = DocxService.replace_tags_in_doc(
            tags_repetidas_bytes,
            replacements_tags_repetidas
        )

        # Extrai texto
        paragraphs_text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Conta ocorrências da substituição
        count = paragraphs_text.count('VALOR_SUBSTITUIDO')
        assert count >= 3  # Deve ter pelo menos 3 substituições

        # Verifica que a tag foi removida
        assert '{TAG_REPETIDA}' not in paragraphs_text

    def test_replace_tags_empty_replacements(self, lista_vendas_bytes):
        """Deve retornar documento inalterado com replacements vazios"""
        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            {}
        )

        assert isinstance(result_doc, DocumentType)

        # Documento deve ainda conter as tags originais
        text = '\n'.join([p.text for p in result_doc.paragraphs])
        assert '{NOME_CLIENTE}' in text
        assert '{DATA_VENDA}' in text
        assert '{VALOR}' in text

    def test_replace_tags_partial_replacements(self, lista_vendas_bytes):
        """Deve substituir apenas tags especificadas"""
        partial_replacements = {
            'NOME_CLIENTE': 'Maria Santos'
            # Não inclui DATA_VENDA nem VALOR
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            partial_replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Tag substituída
        assert 'Maria Santos' in text
        assert '{NOME_CLIENTE}' not in text

        # Tags não substituídas permanecem
        assert '{DATA_VENDA}' in text
        assert '{VALOR}' in text

    def test_replace_tags_case_sensitivity(self, lista_vendas_bytes):
        """Deve respeitar case sensitivity das tags"""
        # Tags em minúscula (não devem substituir)
        replacements = {
            'nome_cliente': 'Teste Minúsculo',
            'NOME_CLIENTE': 'Teste Correto'
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Deve usar apenas a tag em maiúscula
        assert 'Teste Correto' in text

    def test_replace_tags_special_characters(self, lista_vendas_bytes):
        """Deve substituir tags com caracteres especiais no valor"""
        replacements = {
            'NOME_CLIENTE': 'José María Öztürk',
            'DATA_VENDA': '25/12/2025',
            'VALOR': 'R$ 1.234.567,89'
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])

        assert 'José María Öztürk' in text
        assert 'R$ 1.234.567,89' in text

    def test_replace_tags_numeric_values(self, lista_vendas_bytes):
        """Deve converter valores numéricos para string"""
        replacements = {
            'NOME_CLIENTE': 'Cliente 123',
            'DATA_VENDA': 20251206,  # número
            'VALOR': 1500.50  # float
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Valores devem ser convertidos para string
        assert '20251206' in text or '2025' in text
        assert '1500' in text

    def test_replace_tags_in_tables(self, with_tags_bytes, replacements_with_tags):
        """Deve substituir tags dentro de tabelas"""
        result_doc = DocxService.replace_tags_in_doc(
            with_tags_bytes,
            replacements_with_tags
        )

        # Extrai texto das tabelas
        tables_text = ''
        for table in result_doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    tables_text += cell.text + ' '

        # Verifica substituições em tabelas
        assert 'DOC2PDF API' in tables_text
        assert 'Concluído' in tables_text
        assert '{PROJETO}' not in tables_text
        assert '{STATUS}' not in tables_text

    def test_returns_document_object(self, lista_vendas_bytes, replacements_lista_vendas):
        """Deve retornar objeto Document do python-docx"""
        result = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements_lista_vendas
        )

        assert isinstance(result, Document)
        assert hasattr(result, 'paragraphs')
        assert hasattr(result, 'tables')
        assert hasattr(result, 'save')

    def test_document_saveable(self, lista_vendas_bytes, replacements_lista_vendas, temp_dir):
        """Documento resultante deve ser salvável"""
        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements_lista_vendas
        )

        # Salva em arquivo temporário
        output_path = temp_dir / 'output.docx'
        result_doc.save(str(output_path))

        # Verifica que o arquivo foi criado
        assert output_path.exists()
        assert output_path.stat().st_size > 0

        # Verifica que o arquivo é um DOCX válido
        with open(output_path, 'rb') as f:
            content = f.read()
            assert content.startswith(b'PK')  # ZIP header


class TestDocxServiceEdgeCases:
    """Testes de casos extremos"""

    def test_replace_tags_empty_document(self, empty_doc_bytes):
        """Deve processar documento vazio sem erros"""
        result_doc = DocxService.replace_tags_in_doc(
            empty_doc_bytes,
            {'TAG': 'valor'}
        )

        assert isinstance(result_doc, DocumentType)

    def test_replace_tags_nonexistent_tag(self, lista_vendas_bytes):
        """Deve ignorar tags que não existem no documento"""
        replacements = {
            'TAG_INEXISTENTE': 'Não deve dar erro',
            'OUTRA_TAG': 'Também não'
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        assert isinstance(result_doc, DocumentType)

    def test_replace_tags_very_long_value(self, lista_vendas_bytes):
        """Deve substituir com valores muito longos"""
        long_value = 'X' * 10000
        replacements = {
            'NOME_CLIENTE': long_value
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])
        assert long_value in text

    def test_replace_tags_empty_value(self, lista_vendas_bytes):
        """Deve substituir com string vazia"""
        replacements = {
            'NOME_CLIENTE': '',
            'DATA_VENDA': '',
            'VALOR': ''
        }

        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements
        )

        text = '\n'.join([p.text for p in result_doc.paragraphs])

        # Tags devem ser substituídas (removidas)
        assert '{NOME_CLIENTE}' not in text
        assert '{DATA_VENDA}' not in text
        assert '{VALOR}' not in text


class TestDocxServiceIntegration:
    """Testes de integração do serviço DOCX"""

    def test_full_workflow_lista_vendas(self, lista_vendas_bytes, replacements_lista_vendas, temp_dir):
        """Testa fluxo completo: abrir -> substituir -> salvar -> reabrir"""
        # 1. Substituir tags
        result_doc = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            replacements_lista_vendas
        )

        # 2. Salvar
        output_path = temp_dir / 'lista_vendas_processada.docx'
        result_doc.save(str(output_path))

        # 3. Reabrir e verificar
        reopened_doc = Document(str(output_path))
        reopened_text = '\n'.join([p.text for p in reopened_doc.paragraphs])

        # Verifica que as substituições persistiram
        assert 'João Silva' in reopened_text
        assert '06/12/2025' in reopened_text
        assert 'R$ 1.500,00' in reopened_text
        assert '{NOME_CLIENTE}' not in reopened_text

    def test_multiple_replacements_sequential(self, lista_vendas_bytes):
        """Deve permitir múltiplas substituições sequenciais"""
        # Primeira substituição
        doc1 = DocxService.replace_tags_in_doc(
            lista_vendas_bytes,
            {'NOME_CLIENTE': 'Primeira Substituição'}
        )

        # Salva em bytes
        buffer = io.BytesIO()
        doc1.save(buffer)
        doc1_bytes = buffer.getvalue()

        # Segunda substituição (nas tags restantes)
        doc2 = DocxService.replace_tags_in_doc(
            doc1_bytes,
            {'DATA_VENDA': 'Segunda Substituição'}
        )

        text = '\n'.join([p.text for p in doc2.paragraphs])

        # Ambas substituições devem estar presentes
        assert 'Primeira Substituição' in text
        assert 'Segunda Substituição' in text
