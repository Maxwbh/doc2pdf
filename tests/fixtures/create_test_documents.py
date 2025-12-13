"""
Script para criar documentos DOCX de teste

Este script cria os documentos de teste necessários para os testes unitários
e de integração da API DOC2PDF.

Autor: Maxwell da Silva Oliveira - M&S do Brasil LTDA
Versão: 1.5.2
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

# Diretório de destino
FIXTURES_DIR = os.path.dirname(os.path.abspath(__file__))
DOCUMENTS_DIR = os.path.join(FIXTURES_DIR, 'documents')

# Garante que o diretório existe
os.makedirs(DOCUMENTS_DIR, exist_ok=True)


def create_lista_vendas():
    """Cria documento ListaVendas.docx com tags para substituição"""
    doc = Document()

    # Título
    title = doc.add_paragraph()
    title_run = title.add_run('Listagem de Vendas')
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Separador
    doc.add_paragraph('-' * 40)

    # Campos com tags
    doc.add_paragraph(f'Cliente: {{NOME_CLIENTE}}')
    doc.add_paragraph(f'Data da Venda: {{DATA_VENDA}}')
    doc.add_paragraph(f'Valor: {{VALOR}}')

    # Separador final
    doc.add_paragraph('-' * 40)

    # Salva
    filepath = os.path.join(DOCUMENTS_DIR, 'ListaVendas.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


def create_simple_document():
    """Cria documento simples sem tags"""
    doc = Document()

    doc.add_heading('Documento Simples', 0)
    doc.add_paragraph('Este é um documento de teste simples.')
    doc.add_paragraph('Não contém tags para substituição.')
    doc.add_paragraph('Apenas texto puro para validar conversão básica.')

    filepath = os.path.join(DOCUMENTS_DIR, 'simple.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


def create_with_tags():
    """Cria documento com tags em diferentes locais"""
    doc = Document()

    # Título com tag
    doc.add_heading('Relatório de {MES}/{ANO}', 0)

    # Parágrafos com tags
    doc.add_paragraph('Empresa: {EMPRESA}')
    doc.add_paragraph('Responsável: {RESPONSAVEL}')

    # Tabela com tags
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Light Grid Accent 1'

    table.rows[0].cells[0].text = 'Campo'
    table.rows[0].cells[1].text = 'Valor'
    table.rows[1].cells[0].text = 'Projeto'
    table.rows[1].cells[1].text = '{PROJETO}'
    table.rows[2].cells[0].text = 'Status'
    table.rows[2].cells[1].text = '{STATUS}'

    filepath = os.path.join(DOCUMENTS_DIR, 'with_tags.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


def create_complex_document():
    """Cria documento complexo com múltiplos elementos"""
    doc = Document()

    # Título
    doc.add_heading('Documento Complexo - {TITULO}', 0)

    # Seção 1
    doc.add_heading('1. Informações Gerais', 1)
    doc.add_paragraph(f'Cliente: {{CLIENTE}}')
    doc.add_paragraph(f'Projeto: {{PROJETO}}')
    doc.add_paragraph(f'Data: {{DATA}}')

    # Seção 2 - Tabela
    doc.add_heading('2. Detalhes Financeiros', 1)
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Medium Grid 1 Accent 1'

    # Header
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Item'
    hdr_cells[1].text = 'Quantidade'
    hdr_cells[2].text = 'Valor'

    # Dados com tags
    row1 = table.rows[1].cells
    row1[0].text = '{ITEM1}'
    row1[1].text = '{QTD1}'
    row1[2].text = '{VALOR1}'

    row2 = table.rows[2].cells
    row2[0].text = '{ITEM2}'
    row2[1].text = '{QTD2}'
    row2[2].text = '{VALOR2}'

    row3 = table.rows[3].cells
    row3[0].text = 'TOTAL'
    row3[1].text = ''
    row3[2].text = '{TOTAL}'

    # Seção 3 - Lista
    doc.add_heading('3. Observações', 1)
    doc.add_paragraph('Observação 1: {OBS1}', style='List Bullet')
    doc.add_paragraph('Observação 2: {OBS2}', style='List Bullet')
    doc.add_paragraph('Observação 3: {OBS3}', style='List Bullet')

    # Rodapé
    section = doc.sections[0]
    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = 'Gerado em: {DATA_GERACAO} | Versão: {VERSAO}'

    filepath = os.path.join(DOCUMENTS_DIR, 'complex.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


def create_empty_document():
    """Cria documento vazio para testes de erro"""
    doc = Document()

    filepath = os.path.join(DOCUMENTS_DIR, 'empty.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


def create_tags_repetidas():
    """Cria documento com tags repetidas"""
    doc = Document()

    doc.add_heading('Teste de Tags Repetidas', 0)

    # Tag repetida em parágrafos
    doc.add_paragraph('Primeira ocorrência de {TAG_REPETIDA}')
    doc.add_paragraph('Segunda ocorrência de {TAG_REPETIDA}')
    doc.add_paragraph('Terceira ocorrência de {TAG_REPETIDA}')

    # Tag repetida em tabela
    table = doc.add_table(rows=3, cols=1)
    table.rows[0].cells[0].text = '{TAG_REPETIDA}'
    table.rows[1].cells[0].text = 'Texto normal'
    table.rows[2].cells[0].text = '{TAG_REPETIDA}'

    filepath = os.path.join(DOCUMENTS_DIR, 'tags_repetidas.docx')
    doc.save(filepath)
    print(f'✅ Criado: {filepath}')
    return filepath


if __name__ == '__main__':
    print('🔨 Criando documentos de teste...\n')

    create_lista_vendas()
    create_simple_document()
    create_with_tags()
    create_complex_document()
    create_empty_document()
    create_tags_repetidas()

    print(f'\n✅ Todos os documentos de teste foram criados em: {DOCUMENTS_DIR}')
    print(f'\n📊 Total de arquivos: 6')
